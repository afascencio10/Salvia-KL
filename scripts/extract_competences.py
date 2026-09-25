"""Extrae el catálogo de competencias (acciones institucionales) del consolidado
de rutas de SALVIA hacia src/salvia/config/competences.json.

Uso de un solo tiro: se corre cuando el equipo entrega una versión nueva del Excel.
El JSON generado es el que se versiona en el repo; este script no se importa
desde el backend.

Hoja de origen: "4.EntidadesAcciones"
  A Sector | B Entidad | C Alcance territorial | D Resumen de acciones
  E Acciones Institucionales | F Tiempos de respuesta | G Sustento Normativo

Sector/Entidad/Alcance/Resumen vienen en celdas combinadas que cubren varias
acciones, así que se resuelven leyendo los rangos combinados (el valor real vive
en la esquina superior izquierda). NO se rellena hacia abajo fuera de esos
rangos: un tiempo de respuesta o un sustento normativo ausente se deja vacío en
vez de heredar el de la fila anterior, que sería inventar sustento legal.
"""

import hashlib
import json
import re
import sys
import unicodedata
from pathlib import Path

import openpyxl

SHEET = "4.EntidadesAcciones"
FIRST_DATA_ROW = 3  # 1 = título, 2 = encabezados
COL_SECTOR, COL_ENTITY, COL_SCOPE, COL_SUMMARY = 1, 2, 3, 4
COL_ACTION, COL_RESPONSE_TIME, COL_LEGAL_BASIS = 5, 6, 7


def slug(text, max_len=40):
    """Convierte un nombre de entidad en un código estable: sin tildes, snake_case.

    Se corta a max_len y se cierra con un sufijo derivado del nombre completo. El
    corte solo no basta: el Excel trae "Instituciones Prestadoras de Servicios de
    Salud (IPS)" y la misma con el paréntesis largo de clínicas y hospitales, que
    truncadas quedan idénticas. El sufijo depende del nombre, no de la posición de
    la fila, así que reordenar el Excel no reasigna códigos.
    """
    plain = unicodedata.normalize("NFKD", text)
    plain = "".join(c for c in plain if not unicodedata.combining(c))
    plain = re.sub(r"[^a-zA-Z0-9]+", "_", plain).strip("_").lower()
    digest = hashlib.sha1(text.encode("utf-8")).hexdigest()[:4]
    return f"{plain[:max_len].strip('_')}_{digest}"


def clean(value):
    """Texto de celda sin saltos de línea ni espacios repetidos."""
    if value is None:
        return ""
    return re.sub(r"\s+", " ", str(value)).strip()


def merged_lookup(ws):
    """(fila, columna) -> valor, resolviendo los rangos de celdas combinadas."""
    resolved = {}
    for rng in ws.merged_cells.ranges:
        top_left = ws.cell(rng.min_row, rng.min_col).value
        for row in range(rng.min_row, rng.max_row + 1):
            for col in range(rng.min_col, rng.max_col + 1):
                resolved[(row, col)] = top_left
    return resolved


def extract(xlsx_path):
    workbook = openpyxl.load_workbook(xlsx_path, data_only=True)
    ws = workbook[SHEET]
    resolved = merged_lookup(ws)

    def cell(row, col):
        if (row, col) in resolved:
            return clean(resolved[(row, col)])
        return clean(ws.cell(row, col).value)

    kinds = {}
    order = []
    for row in range(FIRST_DATA_ROW, ws.max_row + 1):
        action = cell(row, COL_ACTION)
        if not action:
            continue

        sector = cell(row, COL_SECTOR).upper()
        entity = cell(row, COL_ENTITY)
        if not sector or not entity:
            raise SystemExit(f"fila {row}: acción sin sector o sin entidad — revisar el Excel")

        code = slug(entity)
        if code not in kinds:
            kinds[code] = {"code": code, "sector": sector, "name": entity, "competences": []}
            order.append(code)
        elif kinds[code]["name"] != entity:
            raise SystemExit(
                f"fila {row}: el código '{code}' choca entre '{kinds[code]['name']}' y '{entity}'"
            )

        # El código de la competencia se deriva del texto de la acción, no de su
        # posición: si el equipo reordena filas en el Excel, las respuestas ya
        # guardadas en las barreras siguen apuntando a la misma competencia.
        digest = hashlib.sha1(action.encode("utf-8")).hexdigest()[:6]
        competence_code = f"{code}_{digest}"

        # Una misma acción puede venir repartida en varias filas cuando tiene más
        # de un sustento normativo (ej. GAULA "Negociar en crisis de
        # atrincheramiento", filas 223-224). Se unifica en una sola competencia
        # juntando los sustentos: en el formulario debe aparecer una opción, no dos
        # idénticas.
        existing = next(
            (c for c in kinds[code]["competences"] if c["code"] == competence_code), None
        )
        if existing is not None:
            for field in ("responseTime", "legalBasis"):
                incoming = cell(row, COL_RESPONSE_TIME if field == "responseTime" else COL_LEGAL_BASIS)
                parts = [p for p in existing[field].split(" | ") if p]
                if incoming and incoming not in parts:
                    parts.append(incoming)
                existing[field] = " | ".join(parts)
            existing["sourceRows"].append(row)
            continue

        kinds[code]["competences"].append(
            {
                "code": competence_code,
                "action": action,
                "scope": cell(row, COL_SCOPE),
                "summary": cell(row, COL_SUMMARY),
                "responseTime": cell(row, COL_RESPONSE_TIME),
                "legalBasis": cell(row, COL_LEGAL_BASIS),
                "sourceRows": [row],
            }
        )

    return {
        "source": f"{Path(xlsx_path).name} — hoja {SHEET}",
        "kinds": [kinds[code] for code in order],
    }


def check(catalog):
    """Valida el catálogo extraído. Cualquier fallo aborta sin escribir el JSON."""
    kinds = catalog["kinds"]
    competences = [c for k in kinds for c in k["competences"]]

    codes = [c["code"] for c in competences]
    duplicated = {c for c in codes if codes.count(c) > 1}
    if duplicated:
        raise SystemExit(f"códigos de competencia repetidos: {sorted(duplicated)}")

    empty = [k["code"] for k in kinds if not k["competences"]]
    if empty:
        raise SystemExit(f"entidades sin competencias: {empty}")

    merged = [c for c in competences if len(c["sourceRows"]) > 1]
    source_rows = sum(len(c["sourceRows"]) for c in competences)

    print(f"entidades: {len(kinds)}")
    print(f"filas de acción leídas: {source_rows}")
    print(f"competencias: {len(competences)}")
    for c in merged:
        print(f"  unificada desde filas {c['sourceRows']}: {c['action'][:60]}")
    print(f"  con tiempo de respuesta: {sum(1 for c in competences if c['responseTime'])}")
    print(f"  con sustento normativo:  {sum(1 for c in competences if c['legalBasis'])}")
    for sector in dict.fromkeys(k["sector"] for k in kinds):
        entities = [k for k in kinds if k["sector"] == sector]
        total = sum(len(k["competences"]) for k in entities)
        print(f"  {sector:22} {len(entities):2} entidades / {total:3} competencias")


def main():
    if len(sys.argv) != 3:
        raise SystemExit("uso: extract_competences.py <origen.xlsx> <destino.json>")
    catalog = extract(sys.argv[1])
    check(catalog)
    destination = Path(sys.argv[2])
    destination.write_text(
        json.dumps(catalog, ensure_ascii=False, indent=1) + "\n", encoding="utf-8"
    )
    print(f"\nescrito: {destination} ({destination.stat().st_size / 1024:.0f} KB)")


if __name__ == "__main__":
    main()
